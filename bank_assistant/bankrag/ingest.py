"""Text extraction (PDF, Word, Excel, CSV, text) and structure-aware chunking."""
import csv
import io
import re
from dataclasses import dataclass, field

from .textproc import clean_text

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}


@dataclass
class Unit:
    """One line / paragraph / table row of a document."""
    text: str
    page: int | None = None
    level: int = 0  # 0 = body text, 1 = chapter-like heading, 2 = article-like heading


@dataclass
class Chunk:
    text: str
    page_start: int | None
    page_end: int | None
    heading: str = ""


@dataclass
class Extraction:
    units: list = field(default_factory=list)
    pages: int = 0
    warning: str = ""


_LEVEL1 = re.compile(
    r"^(الباب|الفصل|القسم|الجزء|الملحق|titre|chapitre|partie|section|annexe|chapter|part)\b",
    re.IGNORECASE,
)
_LEVEL2 = re.compile(
    r"^((المادة|مادة|البند)\s+\S+|(article|art\.)\s*\d+|(article)\s+(premier|1er|unique))",
    re.IGNORECASE,
)


def heading_level(line):
    if len(line) > 150:
        return 0
    if _LEVEL1.match(line):
        return 1
    if _LEVEL2.match(line):
        return 2
    return 0


def _lines_to_units(text, page=None):
    units = []
    for line in clean_text(text).split("\n"):
        line = line.strip()
        if line:
            units.append(Unit(line, page, heading_level(line)))
    return units


# --------------------------------------------------------------------------- PDF + bidi

_RTL_CHAR = re.compile("[\u0590-\u065F\u066A-\u06EF\u06FA-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFE]")
# Latin letters and digits; European and Arabic-Indic digits are both laid out left-to-right.
_LTR_CHAR = re.compile("[A-Za-z0-9\u00C0-\u024F\u0660-\u0669\u06F0-\u06F9]")
_JOINERS = set(" .,:/-'’_")
_MIRROR = str.maketrans("()[]{}<>«»", ")(][}{><»«")


def _is_ltr(glyph):
    return bool(_LTR_CHAR.search(glyph)) and not _RTL_CHAR.search(glyph)


def _is_rtl(glyph):
    return bool(_RTL_CHAR.search(glyph))


def _flip_runs(seq, member):
    """Reverse every maximal run of `member` glyphs (joined by spaces/punctuation)."""
    out, i, n = [], 0, len(seq)
    while i < n:
        if member(seq[i]):
            j = i + 1
            while j < n:
                if member(seq[j]):
                    j += 1
                    continue
                if seq[j] in _JOINERS:
                    k = j
                    while k < n and seq[k] in _JOINERS:
                        k += 1
                    if k < n and member(seq[k]):
                        j = k
                        continue
                break
            out.extend(reversed(seq[i:j]))
            i = j
        else:
            out.append(seq[i])
            i += 1
    return out


def visual_to_logical(glyphs, rtl_hint=None):
    """Turn one line of glyph texts in visual (left-to-right) order into logical order.

    Arabic PDFs store glyphs as they are drawn; reading them left-to-right puts
    words backwards and numbers in the wrong place ("الشهري3المادة"). This is a
    compact version of the Unicode bidi algorithm sufficient for Arabic/French text.
    `rtl_hint` (the page's dominant direction) settles lines whose two ends disagree.
    """
    strong = [("R" if _is_rtl(g) else "L") for g in glyphs if _is_rtl(g) or _is_ltr(g)]
    if "R" not in strong:
        return "".join(glyphs)
    if strong[0] == strong[-1]:
        rtl = strong[0] == "R"
    elif rtl_hint is not None:
        rtl = rtl_hint
    else:
        rtl = strong.count("R") >= strong.count("L")
    if rtl:
        # In Arabic context "%" after a number is its own right-to-left item ("30%" is drawn "%30").
        seq = _flip_runs(list(reversed(glyphs)), _is_ltr)
        # brackets drawn in a right-to-left context are mirrored glyphs
        return "".join(g.translate(_MIRROR) if len(g) == 1 else g for g in seq)
    return "".join(_flip_runs(list(glyphs), _is_rtl))


def _pdf_lines_pdfminer(data):
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTChar, LTTextContainer, LTTextLine

    pages = []
    for page in extract_pages(io.BytesIO(data), laparams=LAParams(line_margin=0.4)):
        raw_lines = []
        for box in page:
            if not isinstance(box, LTTextContainer):
                continue
            for line in box:
                if not isinstance(line, LTTextLine):
                    continue
                chars = sorted((c for c in line if isinstance(c, LTChar)), key=lambda c: c.x0)
                glyphs, prev = [], None
                for char in chars:
                    text = char.get_text()
                    if prev is not None and not text.isspace() and not (glyphs and glyphs[-1].isspace()):
                        if char.x0 - prev.x1 > 0.18 * max(char.size, 1):
                            glyphs.append(" ")
                    glyphs.append(text)
                    prev = char
                raw_lines.append(glyphs)
        rtl_count = sum(_is_rtl(g) for glyphs in raw_lines for g in glyphs)
        ltr_count = sum(_is_ltr(g) for glyphs in raw_lines for g in glyphs)
        page_rtl = rtl_count >= ltr_count
        lines = [visual_to_logical(glyphs, page_rtl).strip() for glyphs in raw_lines]
        pages.append("\n".join(line for line in lines if line))
    return pages


def _pdf_lines_pypdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ValueError("encrypted_pdf") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - one bad page must not fail the document
            pages.append("")
    return pages


def extract_pdf(data):
    try:
        pages = _pdf_lines_pdfminer(data)
    except Exception as exc:  # noqa: BLE001 - fall back to the simpler extractor
        if "encrypt" in str(exc).lower() or "password" in str(exc).lower():
            raise ValueError("encrypted_pdf") from exc
        pages = _pdf_lines_pypdf(data)
    result = Extraction(pages=len(pages))
    text_chars = 0
    for number, text in enumerate(pages, start=1):
        text_chars += len(text.strip())
        result.units.extend(_lines_to_units(text, number))
    if result.pages and text_chars < 30 * result.pages:
        result.warning = "scanned_pdf"
    return result


def extract_docx(data):
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(io.BytesIO(data))
    result = Extraction()
    page = 1
    saw_break = False

    def page_breaks(element):
        return len(element.findall(".//" + qn("w:lastRenderedPageBreak"))) + len(
            [br for br in element.findall(".//" + qn("w:br")) if br.get(qn("w:type")) == "page"]
        )

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            breaks = page_breaks(child)
            if breaks:
                saw_break = True
                page += breaks
            paragraph = docx.text.paragraph.Paragraph(child, document)
            text = clean_text(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            level = heading_level(text)
            if style.startswith(("heading 1", "titre 1", "title", "titre")):
                level = 1
            elif style.startswith(("heading", "titre")):
                level = level or 2
            for line in text.split("\n"):
                if line.strip():
                    result.units.append(Unit(line.strip(), page, level))
        elif child.tag == qn("w:tbl"):
            table = docx.table.Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = clean_text(cell.text).replace("\n", " ")
                    if value and (not cells or cells[-1] != value):  # merged cells repeat
                        cells.append(value)
                if cells:
                    result.units.append(Unit(" | ".join(cells), page, 0))
    if not saw_break:
        for unit in result.units:
            unit.page = None
        result.pages = 0
    else:
        result.pages = page
    return result


def _rows_to_units(rows, sheet_number, sheet_name):
    units = []
    if sheet_name:
        units.append(Unit(sheet_name, sheet_number, 1))
    header = None
    for row in rows:
        values = ["" if v is None else clean_text(str(v)).replace("\n", " ") for v in row]
        if not any(values):
            continue
        if header is None:
            header = values
            units.append(Unit(" | ".join(v for v in values if v), sheet_number, 0))
            continue
        parts = []
        for i, value in enumerate(values):
            if not value:
                continue
            name = header[i] if i < len(header) and header[i] else ""
            parts.append(f"{name}: {value}" if name else value)
        units.append(Unit(" | ".join(parts), sheet_number, 0))
    return units


def extract_xlsx(data):
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    result = Extraction()
    for number, sheet in enumerate(workbook.worksheets, start=1):
        result.units.extend(_rows_to_units(sheet.iter_rows(values_only=True), number, sheet.title))
    result.pages = len(workbook.worksheets)
    workbook.close()
    return result


def decode_text(data):
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    for encoding in ("cp1256", "latin-1"):  # Arabic Windows, then Western
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def extract_csv(data):
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return Extraction(units=_rows_to_units(rows, None, ""), pages=0)


def extract_text_file(data):
    return Extraction(units=_lines_to_units(decode_text(data)), pages=0)


def extract(data, ext):
    ext = ext.lower()
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext == ".xlsx":
        return extract_xlsx(data)
    if ext == ".csv":
        return extract_csv(data)
    if ext in (".txt", ".md"):
        return extract_text_file(data)
    raise ValueError("unsupported_type")


# --------------------------------------------------------------------------- chunking

_SENTENCE_END = re.compile(r"(?<=[.!?؟؛;:])\s+")


def _split_long(unit, size):
    """Split a unit longer than `size` on sentence boundaries, then hard-cut."""
    pieces, current = [], ""
    for sentence in _SENTENCE_END.split(unit.text):
        while len(sentence) > size:
            cut = sentence.rfind(" ", 0, size)
            cut = cut if cut > size // 2 else size
            if current:
                pieces.append(current)
                current = ""
            pieces.append(sentence[:cut].strip())
            sentence = sentence[cut:].strip()
        if current and len(current) + 1 + len(sentence) > size:
            pieces.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return [Unit(p, unit.page, unit.level) for p in pieces if p]


def chunk_units(units, size=1200, overlap=200):
    """Group units into chunks of about `size` characters.

    Chunks prefer to start at headings (chapter / article), carry `overlap`
    characters of trailing context otherwise, and remember the heading path
    (e.g. "الفصل الثالث › المادة 12") they belong to.
    """
    expanded = []
    for unit in units:
        expanded.extend(_split_long(unit, size) if len(unit.text) > size else [unit])

    chunks = []
    current = []
    current_len = 0
    fresh = 0  # units added since the last flush (overlap units excluded)
    headings = ["", ""]
    chunk_heading = ""

    def heading_path():
        return " › ".join(h for h in headings if h)

    def flush(keep_overlap):
        nonlocal current, current_len, chunk_heading, fresh
        fresh = 0
        if current:
            text = "\n".join(u.text for u in current).strip()
            pages = [u.page for u in current if u.page is not None]
            if len(text) >= 15:
                chunks.append(Chunk(text, min(pages) if pages else None,
                                    max(pages) if pages else None, chunk_heading))
        tail = []
        if keep_overlap and overlap > 0:
            total = 0
            for unit in reversed(current):
                if total + len(unit.text) > overlap:
                    break
                tail.insert(0, unit)
                total += len(unit.text) + 1
        current = tail
        current_len = sum(len(u.text) + 1 for u in tail)
        chunk_heading = heading_path()

    for unit in expanded:
        if unit.level:
            # a chapter always opens a new chunk; an article does once the chunk has some substance
            if fresh and (unit.level == 1 or current_len > size * 0.35):
                flush(keep_overlap=False)
            elif fresh == 0:
                current, current_len = [], 0  # only overlap from the previous section
            if unit.level == 1:
                headings[0], headings[1] = unit.text[:120], ""
            else:
                headings[1] = unit.text[:120]
            if not current:
                chunk_heading = heading_path()
        if current_len + len(unit.text) + 1 > size and current:
            flush(keep_overlap=True)
        if not current:
            chunk_heading = heading_path()
        current.append(unit)
        current_len += len(unit.text) + 1
        fresh += 1
    flush(keep_overlap=False)
    return chunks
