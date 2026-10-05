import io
import json
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import fake_llm  # noqa: E402
from bankrag import create_app  # noqa: E402
from bankrag.config import Config  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures"


@pytest.fixture(scope="session")
def llm_server():
    server = fake_llm.start()
    yield f"http://127.0.0.1:{server.server_address[1]}/v1"
    server.shutdown()


def make_config(tmp_path, llm_url, **overrides):
    values = dict(DATA_DIR=str(tmp_path / "data"), LLM_BASE_URL=llm_url, LLM_MODEL="fake-hunyuan",
                  LLM_NO_THINK_PREFIX="/no_think", LLM_EXTRA_BODY={}, EMBEDDING_MODEL="",
                  SECRET_KEY="test-secret", CHUNK_SIZE=700, CHUNK_OVERLAP=120, OCR_ENGINE="off")
    values.update(overrides)
    return Config(**values)


@pytest.fixture
def app(tmp_path, llm_server):
    application = create_app(make_config(tmp_path, llm_server))
    application.config["TESTING"] = True
    return application


def csrf_from(html):
    return re.search(r'name="csrf-token" content="([^"]+)"', html).group(1)


class Client:
    """Test client that keeps the CSRF token of the current session."""

    def __init__(self, app):
        self.app = app
        self.http = app.test_client()
        self.csrf = None

    def page(self, path):
        response = self.http.get(path)
        if response.status_code == 200 and b"csrf-token" in response.data:
            self.csrf = csrf_from(response.get_data(as_text=True))
        return response

    def form(self, path, data):
        if self.csrf is None:
            self.page(path)
        return self.http.post(path, data=dict(data, csrf_token=self.csrf))

    def api(self, method, path, payload=None, **kwargs):
        headers = {"X-CSRF-Token": self.csrf or ""}
        if payload is not None:
            kwargs["json"] = payload
        return self.http.open(path, method=method, headers=headers, **kwargs)

    def login(self, username, password):
        self.page("/login")
        response = self.form("/login", {"username": username, "password": password})
        self.page("/")
        return response

    def upload(self, name, data, category=""):
        return self.http.post("/api/documents", headers={"X-CSRF-Token": self.csrf},
                              data={"files": (io.BytesIO(data), name), "category": category},
                              content_type="multipart/form-data")

    def ask(self, question, history=None, doc_ids=None, mode=None):
        payload = {"question": question, "history": history or [], "doc_ids": doc_ids}
        if mode:
            payload["mode"] = mode
        return self.events(self.api("POST", "/api/ask", payload))

    def summary(self, doc_id, lang="ar", refresh=False, follow=False):
        return self.events(self.api("POST", f"/api/documents/{doc_id}/summary",
                                    {"lang": lang, "refresh": refresh, "follow": follow}))

    @staticmethod
    def events(response):
        if response.status_code != 200:
            return response, []
        return response, [json.loads(line) for line in response.get_data(as_text=True).splitlines() if line.strip()]


def create_admin(app, username="admin", password="admin-pass-123"):
    client = Client(app)
    client.page("/setup")
    response = client.form("/setup", {"code": app.config["SETUP_CODE"], "display_name": "Admin",
                                      "username": username, "password": password, "confirm": password})
    assert response.status_code == 302
    return username, password


@pytest.fixture
def admin_client(app):
    username, password = create_admin(app)
    client = Client(app)
    client.login(username, password)
    return client


def sample_docx_fr():
    import docx

    document = docx.Document()
    document.add_heading("Conditions du crédit à la consommation", level=1)
    document.add_paragraph("Article 1 : Le crédit à la consommation finance l'achat de biens durables "
                           "de production nationale.")
    document.add_paragraph("Article 2 : La durée maximale du crédit à la consommation est de 60 mois.")
    document.add_paragraph("Article 3 : Le taux d'intérêt annuel est fixé à 8,25 % hors taxes.")
    table = document.add_table(rows=3, cols=2)
    for row, (a, b) in zip(table.rows, [("Montant", "Taux"), ("Jusqu'à 500 000 DA", "8,25 %"),
                                        ("Plus de 500 000 DA", "7,75 %")]):
        row.cells[0].text, row.cells[1].text = a, b
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_xlsx():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Tarifs"
    sheet.append(["Produit", "Frais mensuels", "Plafond retrait"])
    sheet.append(["Carte CIB Classic", "100 DA", "40 000 DA"])
    sheet.append(["Carte CIB Gold", "250 DA", "100 000 DA"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
